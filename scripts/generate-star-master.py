"""Generate STAR greenhouse TypeScript rows + SQL seed from VegPro Excel.

Reads Greenhouse Bay Mapping Star (1).xlsx Final sheet.
Merged area / columns-per-bay cells apply to every bay in the merge.
Half-columns (e.g. 5.5) are ceiled so the heat map always has a whole box.
"""
from __future__ import annotations

import math
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

XLSX = Path(r"C:\Users\PC\Downloads\Greenhouse Bay Mapping Star (1).xlsx")
ROOT = Path(__file__).resolve().parents[1]
TS_PATH = ROOT / "src/lib/star-greenhouse-rows.ts"
SQL_PATH = ROOT / "supabase/migrations/019_star_gh_columns.sql"

ALIASES = {
    "redcalypso": "Red Calypso",
    "red calypso": "Red Calypso",
    "simply orange": "Simply Orange",
    "simplyorange": "Simply Orange",
    "fuschiana": "Fuschiana",
    "athena": "Athena",
    "confidential": "Confidential",
    "madamred": "Madam Red",
    "madam red": "Madam Red",
    "millionreason": "Million Reason",
    "simply pink": "Simply Pink",
    "simplypink": "Simply Pink",
    "moonwalk": "Moonwalk",
    "revolution": "Revolution",
    "escape": "Escape",
    "inka": "Inka",
    "natures white": "Natures White",
    "orange candy": "Orange Candy",
    "pink arrow": "Pink Arrow",
}


def title_variety(name: str) -> str:
    raw = re.sub(r"\s+", " ", name).strip()
    return ALIASES.get(raw.lower(), raw.title())


def normalize_gh(name: str) -> str:
    name = name.strip().upper().replace(" ", "")
    name = re.sub(r"^STGH0*13", "STGH13", name)
    return name


def ceil_columns(value: float) -> int:
    return max(1, int(math.ceil(value - 1e-9)))


def load_existing_ids() -> dict[str, str]:
    if not TS_PATH.exists():
        return {}
    text = TS_PATH.read_text(encoding="utf-8")
    return {name: gid for gid, name in re.findall(r'id: "([^"]+)", name: "([^"]+)"', text)}


def next_id(existing: dict[str, str], index: int) -> str:
    nums = [int(v.split("-")[-1], 16) for v in existing.values()]
    n = max(nums) + index if nums else index
    return f"beeeee02-0001-4001-8001-{n:012d}"


def load_final():
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Final"]
    filled: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(rng))
        value = ws.cell(min_row, min_col).value
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                filled[(r, c)] = value

    def cell(r: int, c: int):
        return filled.get((r, c), ws.cell(r, c).value)

    gh: OrderedDict[str, dict] = OrderedDict()
    for r in range(2, ws.max_row + 1):
        raw_name = cell(r, 2)
        if not raw_name:
            continue
        name = normalize_gh(str(raw_name))
        area = cell(r, 3)
        bay = cell(r, 4)
        cols = cell(r, 5)
        variety = title_variety(str(cell(r, 6) or ""))
        rec = gh.setdefault(name, {"area": None, "bays": {}, "cols": {}})
        if area not in (None, ""):
            rec["area"] = float(area)
        if bay in (None, ""):
            continue
        b = int(float(bay))
        rec["bays"][b] = variety
        if cols not in (None, ""):
            rec["cols"][b] = float(cols)
    return gh


def fmt_area(area: float) -> str:
    if float(area).is_integer():
        return str(int(area))
    return str(area)


def main():
    gh = load_final()
    existing_ids = load_existing_ids()
    new_count = 0
    rows_out = []
    variety_order: list[str] = []

    for i, (name, rec) in enumerate(gh.items(), start=1):
        bays = sorted(rec["bays"]) or [1]
        bay_max = max(bays)
        last_var = rec["bays"].get(bays[0], "")
        filled_var: dict[int, str] = {}
        last_col = rec["cols"].get(bays[0], 11.0)
        filled_cols: dict[int, int] = {}
        for b in range(1, bay_max + 1):
            if rec["bays"].get(b):
                last_var = rec["bays"][b]
            if last_var:
                filled_var[b] = last_var
            if b in rec["cols"]:
                last_col = rec["cols"][b]
            filled_cols[b] = ceil_columns(last_col)

        varieties: list[str] = []
        for b in range(1, bay_max + 1):
            v = filled_var.get(b)
            if v and v not in varieties:
                varieties.append(v)
        for v in varieties:
            if v not in variety_order:
                variety_order.append(v)

        unique_cols = sorted(set(filled_cols.values()))
        column_max = max(unique_cols) if unique_cols else 11
        columns_by_bay = filled_cols if len(unique_cols) > 1 else None

        gh_id = existing_ids.get(name)
        if not gh_id:
            new_count += 1
            gh_id = next_id(existing_ids, new_count)
            existing_ids[name] = gh_id

        rows_out.append(
            {
                "id": gh_id,
                "name": name,
                "areaSqm": rec["area"] or 4000.0,
                "bayMax": bay_max,
                "columnMax": column_max,
                "columnsByBay": columns_by_bay,
                "varieties": varieties,
                "varietyByBay": filled_var if len(varieties) > 1 else None,
            }
        )

    ts_lines = [
        "/** Auto-generated from VegPro Greenhouse Bay Mapping Star (1).xlsx (Final sheet). */",
        "",
        "export type StarGreenhouseRow = {",
        "  id: string;",
        "  name: string;",
        "  areaSqm: number;",
        "  bayMax: number;",
        "  columnMax: number;",
        "  columnsByBay?: Readonly<Record<number, number>>;",
        "  varieties: readonly string[];",
        "  varietyByBay?: Readonly<Record<number, string>>;",
        "};",
        "",
        "export const STAR_VARIETY_ORDER = [",
        "  " + ", ".join(f'"{v}"' for v in variety_order),
        "] as const;",
        "",
        "export const STAR_GREENHOUSE_ROWS: readonly StarGreenhouseRow[] = [",
    ]
    for r in rows_out:
        vars_js = ", ".join(f'"{v}"' for v in r["varieties"])
        extra = ""
        if r["varietyByBay"]:
            pairs = ", ".join(f'{b}: "{v}"' for b, v in sorted(r["varietyByBay"].items()))
            extra += f", varietyByBay: {{ {pairs} }}"
        if r["columnsByBay"]:
            pairs = ", ".join(f"{b}: {c}" for b, c in sorted(r["columnsByBay"].items()))
            extra += f", columnsByBay: {{ {pairs} }}"
        ts_lines.append(
            f'  {{ id: "{r["id"]}", name: "{r["name"]}", areaSqm: {fmt_area(r["areaSqm"])}, bayMax: {r["bayMax"]}, columnMax: {r["columnMax"]}, varieties: [{vars_js}]{extra} }},'
        )
    ts_lines.append("];")
    ts_lines.append("")
    TS_PATH.write_text("\n".join(ts_lines), encoding="utf-8")

    sql = [
        "-- Star GH area, bay counts, and columns-per-bay from Greenhouse Bay Mapping Star (1).xlsx",
        "",
        "alter table public.greenhouses",
        "  add column if not exists area_sqm numeric(10, 2),",
        "  add column if not exists bay_count smallint,",
        "  add column if not exists column_count smallint;",
        "",
        "update public.greenhouses as g set",
        "  area_sqm = v.area_sqm,",
        "  bay_count = v.bay_count,",
        "  column_count = v.column_count,",
        "  area_ha = round(v.area_sqm / 10000.0, 4)",
        "from (values",
    ]
    values = []
    for r in rows_out:
        values.append(
            f"  ('{r['id']}'::uuid, {fmt_area(r['areaSqm'])}::numeric, {r['bayMax']}::smallint, {r['columnMax']}::smallint)"
        )
    sql.append(",\n".join(values))
    sql.append(") as v(id, area_sqm, bay_count, column_count)")
    sql.append("where g.id = v.id;")
    sql.append("")
    SQL_PATH.write_text("\n".join(sql) + "\n", encoding="utf-8")
    print(f"wrote {len(rows_out)} houses, {len(variety_order)} varieties -> {TS_PATH.name}, {SQL_PATH.name}")


if __name__ == "__main__":
    main()
